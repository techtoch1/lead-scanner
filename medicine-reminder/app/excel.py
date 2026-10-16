"""Read ./data/medicines.xlsx into medicines + "needs review" rows.

Nothing is hardcoded: columns are found by header name (English or Arabic
aliases below). Any row that can't be understood unambiguously becomes a
ReviewRow and never gets a reminder.
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

# canonical field -> accepted header names (compared after normalisation)
HEADERS = {
    "name": ["name", "medicine", "medication", "drug", "اسم الدواء", "الدواء", "الاسم"],
    "dose": ["dose", "dosage", "الجرعة", "الجرعه"],
    "times": ["time", "times", "time(s)", "الوقت", "الاوقات", "وقت"],
    "frequency": ["frequency", "freq", "schedule", "التكرار", "تكرار"],
    "day": ["day", "days", "weekday", "day of week", "اليوم", "الايام"],
    "start_date": ["start date", "start", "first dose", "next date", "next dose", "تاريخ البدء", "تاريخ البداية"],
    "last_date": ["last date", "last dose", "last taken", "اخر جرعة", "تاريخ اخر جرعة"],
    "instruction": ["instruction", "instructions", "التعليمات", "تعليمات", "الارشادات"],
    "purpose": ["purpose", "for", "used for", "indication", "what for", "الغرض", "الاستخدام", "لماذا", "دواعي الاستعمال"],
}

WEEKDAYS = {  # python weekday(): Mon=0
    0: ["monday", "mon", "الاثنين", "الاتنين"],
    1: ["tuesday", "tue", "tues", "الثلاثاء"],
    2: ["wednesday", "wed", "الاربعاء"],
    3: ["thursday", "thu", "thur", "thurs", "الخميس"],
    4: ["friday", "fri", "الجمعه"],
    5: ["saturday", "sat", "السبت"],
    6: ["sunday", "sun", "الاحد"],
}

_AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
_TASHKEEL = re.compile("[ً-ٰٟـ]")
_AR_FOLD = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه", "؛": ";", "،": ","})


def norm(value):
    """Lowercase, ASCII digits, folded Arabic letters, no diacritics, single spaces."""
    if value is None:
        return ""
    text = str(value).translate(_AR_DIGITS)
    text = _TASHKEEL.sub("", text).translate(_AR_FOLD).lower()
    return re.sub(r"[\s_]+", " ", text).strip()


@dataclass
class Medicine:
    row: int
    name: str
    dose: str
    instruction: str
    purpose: str
    kind: str  # daily | weekly | every_n | as_needed
    times: list = field(default_factory=list)  # ["08:00", ...]
    weekdays: list = field(default_factory=list)  # python weekdays, for weekly
    every_n: int | None = None
    anchor: date | None = None  # a date the every_n cycle falls on


@dataclass
class ReviewRow:
    row: int
    name: str
    reasons: list
    raw: dict


@dataclass
class ParseResult:
    medicines: list = field(default_factory=list)
    review: list = field(default_factory=list)
    errors: list = field(default_factory=list)  # file-level problems
    mtime: float = 0.0


# ---------- cell parsers ----------

_TIME_RE = re.compile(r"^(\d{1,2})(?:[:.](\d{2}))?\s*(am|pm|ص|م|صباحا|مساء)?$")


def _fmt(hour, minute):
    return f"{hour:02d}:{minute:02d}"


def parse_times(cell):
    """Return (times, problems). Accepts Excel time cells or text like '8:00, 8pm, 20:30'."""
    if cell is None or (isinstance(cell, str) and not cell.strip()):
        return [], ["missing time"]
    if isinstance(cell, (time, datetime)):
        return [_fmt(cell.hour, cell.minute)], []
    if isinstance(cell, (int, float)):
        if 0 <= cell < 1:  # raw Excel day fraction
            minutes = round(cell * 24 * 60)
            return [_fmt(minutes // 60 % 24, minutes % 60)], []
        return [], [f"time '{cell}' is unclear (use HH:MM)"]
    times, problems = [], []
    parts = re.split(r"[,;\n/]|\s+و\s+|\s+and\s+", norm(cell))
    for part in (p.strip() for p in parts):
        if not part:
            continue
        m = _TIME_RE.match(part)
        if not m or not (m.group(2) or m.group(3)):
            problems.append(f"time '{part}' is unclear (use HH:MM)")
            continue
        hour, minute, mer = int(m.group(1)), int(m.group(2) or 0), m.group(3)
        if mer:
            if not 1 <= hour <= 12:
                problems.append(f"time '{part}' is invalid")
                continue
            pm = mer in ("pm", "م", "مساء")
            hour = hour % 12 + (12 if pm else 0)
        if hour > 23 or minute > 59:
            problems.append(f"time '{part}' is invalid")
            continue
        times.append(_fmt(hour, minute))
    if not times and not problems:
        problems.append("missing time")
    return sorted(set(times)), problems


def parse_date(cell):
    """Return (date|None, problem|None). Text dates: ISO, or DD/MM/YYYY."""
    if cell is None or (isinstance(cell, str) and not cell.strip()):
        return None, None
    if isinstance(cell, datetime):
        return cell.date(), None
    if isinstance(cell, date):
        return cell, None
    text = norm(cell)
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", text)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$", text)
        if not m:
            return None, f"date '{cell}' is unclear (use YYYY-MM-DD)"
        d, mo, y = map(int, m.groups())  # day first
    try:
        return date(y, mo, d), None
    except ValueError:
        return None, f"date '{cell}' is invalid"


def find_weekdays(text):
    found = []
    for wd, names in WEEKDAYS.items():
        if any(re.search(rf"(?<![a-z؀-ۿ]){re.escape(n)}(?![a-z؀-ۿ])", text) for n in names):
            found.append(wd)
    return found


AS_NEEDED = ["as needed", "prn", "when needed", "if needed", "عند الحاجه", "اذا لزم", "حسب الحاجه", "عند اللزوم"]
DAILY = ["daily", "every day", "once a day", "everyday", "يوميا", "يومي", "كل يوم"]
WEEKLY = ["weekly", "every week", "once a week", "اسبوعيا", "اسبوعي", "كل اسبوع"]


def classify_frequency(text):
    """Return (kind, every_n, problems). kind: daily|weekly|every_n|as_needed|None."""
    if not text:
        return None, None, ["missing frequency"]
    if any(k in text for k in AS_NEEDED):
        return "as_needed", None, []
    m = re.search(r"every\s+(\d+)\s*(day|days|week|weeks)\b", text) or re.search(
        r"كل\s*(\d+)\s*(يوم|ايام|اسبوع|اسابيع)", text
    )
    if m:
        n = int(m.group(1))
        if m.group(2) in ("week", "weeks", "اسبوع", "اسابيع"):
            n *= 7
        if n < 1:
            return None, None, [f"frequency '{text}' is unclear"]
        return "every_n", n, []
    if any(k in text for k in DAILY):
        return "daily", None, []
    if any(k in text for k in WEEKLY):
        return "weekly", None, []
    return None, None, [f"frequency '{text}' is unclear"]


# ---------- sheet reading ----------

def _column_map(header_cells):
    alias = {norm(a): canon for canon, names in HEADERS.items() for a in names}
    cols = {}
    for idx, cell in enumerate(header_cells):
        canon = alias.get(norm(cell))
        if canon and canon not in cols:
            cols[canon] = idx
    return cols


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_row(row_no, rec):
    """rec: canonical field -> raw cell. Returns Medicine or ReviewRow."""
    name = _text(rec.get("name"))
    raw = {k: _text(v) for k, v in rec.items() if v not in (None, "")}
    problems = []
    if not name:
        problems.append("missing medicine name")

    freq_text = norm(rec.get("frequency"))
    day_text = norm(rec.get("day"))
    kind, n, fprob = classify_frequency(freq_text)
    problems += fprob
    if not freq_text and day_text:
        problems.append("frequency is empty (a day is given — is it weekly?)")

    times, anchor, weekdays = [], None, []
    if kind and kind != "as_needed":
        times, tprob = parse_times(rec.get("times"))
        problems += tprob

    if kind == "weekly":
        weekdays = find_weekdays(day_text) or find_weekdays(freq_text)
        if not weekdays:
            problems.append("weekly medicine needs a day of the week")
    elif kind == "every_n":
        start, p1 = parse_date(rec.get("start_date"))
        last, p2 = parse_date(rec.get("last_date"))
        problems += [p for p in (p1, p2) if p]
        if start and last:
            problems.append("both start date and last date given — keep one")
        elif start:
            anchor = start
        elif last:
            anchor = last + timedelta(days=n)
        elif not (p1 or p2):
            problems.append(f"every {n} days needs a start date or last-dose date")

    if problems:
        return ReviewRow(row_no, name, problems, raw)

    if kind == "every_n" and n == 1:
        kind, anchor = "daily", None
    elif kind == "every_n" and n == 7:
        kind, weekdays, anchor = "weekly", [anchor.weekday()], None
        n = None
    return Medicine(
        row=row_no, name=name, dose=_text(rec.get("dose")), instruction=_text(rec.get("instruction")),
        purpose=_text(rec.get("purpose")), kind=kind, times=times, weekdays=weekdays,
        every_n=n if kind == "every_n" else None, anchor=anchor,
    )


def load_medicines(path):
    path = Path(path)
    result = ParseResult()
    if not path.exists():
        result.errors.append(f"Excel file not found: {path.name} (put it in the data folder)")
        return result
    result.mtime = path.stat().st_mtime
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # corrupt / not xlsx / locked
        result.errors.append(f"Could not open {path.name}: {type(exc).__name__}")
        return result
    try:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        result.errors.append("The Excel sheet is empty")
        return result
    cols = _column_map(rows[0])
    missing = [c for c in ("name", "times", "frequency") if c not in cols]
    if missing:
        result.errors.append("Missing required column(s) in the header row: " + ", ".join(missing))
        return result
    for offset, cells in enumerate(rows[1:], start=2):
        rec = {c: (cells[i] if i < len(cells) else None) for c, i in cols.items()}
        if all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in rec.values()):
            continue
        item = parse_row(offset, rec)
        (result.medicines if isinstance(item, Medicine) else result.review).append(item)
    return result
